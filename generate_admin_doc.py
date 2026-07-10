#!/usr/bin/env python3
"""Generate an editable Word (.docx) admin document from the Resources Planner.

Contains two sections:
  1. Module Deliveries — one row per module with period(s), tutor(s), moderator(s)
  2. Programme Leaders — staff with Programme Management roles

Usage:
    python3 generate_admin_doc.py MASTER.xlsm [-o admin_doc.docx]
"""

import argparse
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

from generate_allocation_pdfs import (
    SHEET, find_sections,
)
from generate_admin_sheet import (
    build_module_tutors, build_module_period_map, build_module_title_map,
    build_module_moderators, build_programme_leaders,
)

TEAL      = RGBColor(0x1F, 0x3E, 0x64)
TEAL_LIGHT = RGBColor(0xE8, 0xEE, 0xF0)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text) if text else "")
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    cell.paragraphs[0].paragraph_format.space_after  = Pt(2)


def add_header_row(table, headers, col_widths_cm):
    row = table.rows[0]
    for i, (h, w) in enumerate(zip(headers, col_widths_cm)):
        cell = row.cells[i]
        set_cell_bg(cell, "1F3E64")
        cell_text(cell, h, bold=True, color=WHITE, size=10)
        cell.width = Cm(w)


def add_data_row(table, values, alt=False):
    row = table.add_row()
    for i, v in enumerate(values):
        cell = row.cells[i]
        if alt:
            set_cell_bg(cell, "EEF2F8")
        cell_text(cell, v, size=10)
    return row


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = TEAL
    return p


def build_doc(tutors, periods, titles, moderators, pm_rows):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2)
    section.right_margin  = Cm(2)

    # Title
    title = doc.add_heading("School of Computer Science — Workload Admin", 0)
    for run in title.runs:
        run.font.color.rgb = TEAL

    doc.add_paragraph("2026–27 Academic Year").paragraph_format.space_after = Pt(6)

    # ── Section 1: Module Deliveries ────────────────────────────────────────
    heading(doc, "Module Deliveries", level=1)

    col_widths = [2.8, 6.5, 2.8, 5.5, 5.5]
    headers    = ["Module Code", "Module Title", "Period(s)", "Module Tutor(s)", "Moderator(s)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    add_header_row(table, headers, col_widths)

    for i, code in enumerate(sorted(tutors)):
        add_data_row(table, [
            code,
            titles.get(code, ""),
            ", ".join(periods.get(code, [])),
            ", ".join(tutors[code]),
            ", ".join(moderators.get(code, [])),
        ], alt=(i % 2 == 0))

    doc.add_paragraph()

    # ── Section 2: Programme Leaders ────────────────────────────────────────
    heading(doc, "Programme Leaders", level=1)

    col_widths2 = [5.0, 2.5, 5.5, 2.0, 8.0]
    headers2    = ["Staff Name", "Staff ID", "Role", "Hours", "Programme(s)"]
    table2 = doc.add_table(rows=1, cols=len(headers2))
    table2.style = "Table Grid"
    add_header_row(table2, headers2, col_widths2)

    for i, (name, sid, role, hrs) in enumerate(pm_rows):
        add_data_row(table2, [name, sid, role, hrs, ""], alt=(i % 2 == 0))

    return doc


def main():
    parser = argparse.ArgumentParser(description="Generate editable Word admin document")
    parser.add_argument("workbook", help="Resources Planner workbook (.xlsm)")
    parser.add_argument("-o", "--output", default="admin_doc.docx",
                        help="Output .docx path (default: admin_doc.docx)")
    args = parser.parse_args()

    print(f"Reading {args.workbook} …")
    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb[SHEET]
    sec = find_sections(ws)

    tutors     = build_module_tutors(ws, sec)
    periods    = build_module_period_map(ws, sec)
    titles     = build_module_title_map(ws, sec)
    moderators = build_module_moderators(ws, sec)
    pm_rows    = build_programme_leaders(ws, sec)

    doc = build_doc(tutors, periods, titles, moderators, pm_rows)

    out = Path(args.output)
    doc.save(out)
    print(f"Written: {out}")
    print(f"  Module Deliveries: {len(tutors)} modules")
    print(f"  Programme Leaders: {len(pm_rows)} rows")


if __name__ == "__main__":
    main()
