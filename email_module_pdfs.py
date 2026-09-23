#!/usr/bin/env python3
"""Send each module tutor their module summary PDF via macOS Mail.app."""

import csv
import datetime
import subprocess
import sys
from pathlib import Path

import openpyxl
from generate_allocation_pdfs import SHEET, find_sections, list_staff, build_module_tutors, safe_name

EMAIL_CSV = Path(__file__).parent / "staff_emails.csv"
MODULE_PDF_DIR = Path(__file__).parent / "module_pdfs"

EMAIL_SUBJECT = "2026-27 Module Workload Summaries"
_TODAY = datetime.date.today().strftime("%d %B %Y")
EMAIL_BODY = f"""\
Dear {{first_name}},

Please find attached the workload summaries (generated on {_TODAY}) for the module(s) you are listed as module tutor for:

{{module_list}}
{{no_pdf_note}}

These documents are indicative, and are being shared to aid in transparency around the allocation process.

This is a brand new process, so please bear with us as we refine the document generation process.

If you have any questions about what these documents show, please don't hesitate to get in touch with myself or another subject head.

Best regards,

Tom Bartindale
Head of Subject, Computing


"""

# ── Email lookup from CSV ─────────────────────────────────────────────────────

def load_email_map() -> dict[str, str]:
    if not EMAIL_CSV.exists():
        sys.exit(f"Email map not found: {EMAIL_CSV}")
    mapping = {}
    with open(EMAIL_CSV, newline="") as f:
        for row in csv.DictReader(f):
            if row["email"].strip():
                mapping[row["staff_id"].strip()] = row["email"].strip()
    return mapping

# ── Send via Mail.app ─────────────────────────────────────────────────────────

def send_email(to_email: str, subject: str, body: str, pdf_paths: list):
    attachments = "\n".join(
        f'make new attachment with properties {{file name:POSIX file "{str(p.resolve())}"}} at after last paragraph'
        for p in pdf_paths
    )
    script = f"""
tell application "Mail"
    set msg to make new outgoing message with properties {{subject:"{subject}", content:"{body}", visible:false}}
    tell msg
        make new to recipient at end of to recipients with properties {{address:"{to_email}"}}
        {attachments}
    end tell
    send msg
end tell
"""
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True)

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Email module summary PDFs to module tutors via macOS Mail.app")
    parser.add_argument("workbook", help="Resources Planner workbook (.xlsm)")
    parser.add_argument("--pdf-dir", default=str(MODULE_PDF_DIR),
                        help=f"Folder of module PDFs (default: {MODULE_PDF_DIR})")
    parser.add_argument("--dry-run", action="store_true",
                        help="Show what would be sent without sending anything")
    parser.add_argument("--module", metavar="CODE",
                        help="Only send for this module code")
    parser.add_argument("--id", metavar="STAFF_ID",
                        help="Only send to this staff member")
    args = parser.parse_args()

    pdf_dir = Path(args.pdf_dir)
    if not pdf_dir.is_dir():
        sys.exit(f"Directory not found: {pdf_dir}")

    email_map = load_email_map()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb[SHEET]
    sec = find_sections(ws)

    # staff_id -> email already in email_map; build name -> staff_id
    name_to_sid: dict[str, str] = {}
    for col, sid, _full in list_staff(ws):
        forename = ws.cell(row=3, column=col).value or ""
        surname  = ws.cell(row=4, column=col).value or ""
        name = f"{forename} {surname}".strip()
        if name:
            name_to_sid[name] = sid

    # Module title lookup
    from generate_allocation_pdfs import COL_B, COL_TITLE
    ta_start, ta_end = sec["Teaching and Assessment"]
    module_titles: dict[str, str] = {}
    for r in range(ta_start + 2, ta_end + 1):
        code  = ws.cell(row=r, column=COL_B).value
        title = ws.cell(row=r, column=COL_TITLE).value
        if code and title and str(code).strip() not in module_titles:
            module_titles[str(code).strip()] = str(title).strip()

    tutors = build_module_tutors(ws, sec)
    if args.module:
        tutors = {k: v for k, v in tutors.items() if k == args.module}

    # Group by staff member: name -> [(code, pdf_path)]
    from collections import defaultdict
    staff_modules: dict[str, list] = defaultdict(list)
    staff_no_pdf: dict[str, list] = defaultdict(list)
    no_pdf = []
    for code, tutor_entries in sorted(tutors.items()):
        names = sorted({name for name, _hint in tutor_entries})
        pdf_path = pdf_dir / f"{safe_name(code)}_Module.pdf"
        if not pdf_path.exists():
            no_pdf.append(code)
            for name in names:
                staff_no_pdf[name].append(code)
            continue
        for name in names:
            staff_modules[name].append((code, pdf_path))

    if no_pdf:
        print(f"No PDF found for: {', '.join(no_pdf)}\n")

    total_staff = len(set(staff_modules) | set(staff_no_pdf))
    print(f"Sending to {total_staff} staff member(s).{' Dry run — no emails will be sent.' if args.dry_run else ''}\n")

    sent, skipped, failed = [], [], []

    all_staff_names = sorted(set(staff_modules) | set(staff_no_pdf))
    if args.id:
        target_sid = args.id.upper()
        all_staff_names = [n for n in all_staff_names if name_to_sid.get(n) == target_sid]
    for name in all_staff_names:
        modules = staff_modules.get(name, [])
        sid   = name_to_sid.get(name)
        email = email_map.get(sid) if sid else None
        first_name = name.split()[0] if name else "colleague"

        if not email:
            print(f"  SKIP  {name}  — no email in staff_emails.csv")
            skipped.append(name)
            continue

        module_list = "".join(
            f"  - {code}: {module_titles.get(code, '')}\n"
            for code, _ in modules
        )
        missing = staff_no_pdf.get(name, [])
        if missing:
            missing_lines = "".join(f"  - {code}: {module_titles.get(code, '')}\n" for code in missing)
            no_pdf_note = (
                f"You are also listed as module tutor for the following module(s), "
                f"for which a separate summary has not been included:\n\n{missing_lines}\n"
            )
        else:
            no_pdf_note = ""
        body = EMAIL_BODY.format(first_name=first_name, module_list=module_list, no_pdf_note=no_pdf_note)
        codes_str = ", ".join(code for code, _ in modules)
        if missing:
            codes_str += f"  (no PDF: {', '.join(missing)})"

        if args.dry_run:
            print(f"  DRY   {name:<35} → {email}  [{codes_str}]")
            sent.append(name)
            continue

        try:
            pdfs = [p for _, p in modules]
            send_email(email, EMAIL_SUBJECT, body, pdfs)
            print(f"  SENT  {name:<35} → {email}  [{codes_str}]")
            sent.append(name)
        except subprocess.CalledProcessError as exc:
            print(f"  FAIL  {name:<35} → {email}  [{exc.stderr.decode().strip()}]")
            failed.append(name)

    suffix = " (dry run)" if args.dry_run else ""
    print(f"\nDone{suffix}. {len(sent)} sent, {len(skipped)} skipped (no email), {len(failed)} failed.")
    if failed:
        print("Failed:", ", ".join(failed))

if __name__ == "__main__":
    main()
