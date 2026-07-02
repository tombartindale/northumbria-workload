#!/usr/bin/env python3
"""Generate an admin spreadsheet with two sheets:
  1. Module Tutors  — one row per module with period(s) and tutor names
  2. Programme Leaders — staff with Programme Management (non-tutor) roles,
     with a blank 'Programme(s)' column to fill in
"""

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from generate_allocation_pdfs import (
    SHEET, COL_B, COL_TITLE, COL_PERIOD, COL_ACTIVITY,
    find_sections, list_staff, num,
)

PERIOD_ORDER = {"SEM1": 1, "SEM2": 2, "YL": 3, "YLSEM1": 4, "YLSEM2": 5,
                "Semester 3": 6, "SEM3": 6, "YLSEM3": 6}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL    = PatternFill("solid", fgColor="EEF2F8")
THIN_SIDE   = Side(style="thin", color="CCCCCC")
CELL_BORDER = Border(bottom=THIN_SIDE)


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_row(ws, row, cols, alt=False):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        if alt:
            c.fill = ALT_FILL
        c.border = CELL_BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_module_period_map(ws, sec):
    periods = defaultdict(set)
    ta_start, ta_end = sec["Teaching and Assessment"]
    for r in range(ta_start + 2, ta_end + 1):
        code   = ws.cell(row=r, column=COL_B).value
        period = ws.cell(row=r, column=COL_PERIOD).value
        if code and period:
            periods[str(code).strip()].add(str(period).strip())
    return {
        code: sorted(ps, key=lambda p: PERIOD_ORDER.get(p, 99))
        for code, ps in periods.items()
    }


def build_module_title_map(ws, sec):
    ta_start, ta_end = sec["Teaching and Assessment"]
    titles = {}
    for r in range(ta_start + 2, ta_end + 1):
        code  = ws.cell(row=r, column=COL_B).value
        title = ws.cell(row=r, column=COL_TITLE).value
        if code and title and str(code).strip() not in titles:
            titles[str(code).strip()] = str(title).strip()
    return titles


def build_module_tutors(ws, sec):
    staff_names = {}
    for col, _sid, _full in list_staff(ws):
        name = f"{ws.cell(row=3, column=col).value or ''} {ws.cell(row=4, column=col).value or ''}".strip()
        staff_names[col] = name

    tutors = defaultdict(set)
    pm_start, pm_end = sec["Programme Management"]
    for r in range(pm_start, pm_end + 1):
        lab = ws.cell(row=r, column=COL_ACTIVITY).value
        if isinstance(lab, str) and lab.startswith("Module Code"):
            for col, name in staff_names.items():
                if not name:
                    continue
                code = ws.cell(row=r, column=col).value
                hrs  = num(ws.cell(row=r + 1, column=col).value)
                if code and hrs:
                    tutors[str(code).strip()].add(name)
    return {k: sorted(v) for k, v in tutors.items()}


def build_programme_leaders(ws, sec):
    staff_cols = {}
    for col, sid, _full in list_staff(ws):
        name = f"{ws.cell(row=3, column=col).value or ''} {ws.cell(row=4, column=col).value or ''}".strip()
        staff_cols[col] = (name, sid)

    rows = []
    pm_start, pm_end = sec["Programme Management"]
    for r in range(pm_start + 2, pm_end + 1):
        lab    = ws.cell(row=r, column=COL_B).value
        actlab = ws.cell(row=r, column=COL_ACTIVITY).value
        if isinstance(actlab, str) and (actlab.startswith("Module Code") or actlab == "Hours"):
            continue
        if not lab or str(lab).endswith("Total"):
            continue
        role = str(lab).strip()
        for col, (name, sid) in staff_cols.items():
            if not name:
                continue
            hrs = num(ws.cell(row=r, column=col).value)
            if hrs:
                rows.append((name, sid, role, hrs))

    rows.sort(key=lambda x: (x[2], x[0]))
    return rows


def write_module_tutors_sheet(wb, tutors, periods, titles):
    ws = wb.create_sheet("Module Tutors")
    headers = ["Module Code", "Module Title", "Period(s)", "Module Tutor(s)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 22

    for row, code in enumerate(sorted(tutors), 2):
        ws.cell(row=row, column=1, value=code)
        ws.cell(row=row, column=2, value=titles.get(code, ""))
        ws.cell(row=row, column=3, value=", ".join(periods.get(code, [])))
        ws.cell(row=row, column=4, value=", ".join(tutors[code]))
        style_row(ws, row, len(headers), alt=(row % 2 == 0))
        ws.row_dimensions[row].height = 18

    set_col_widths(ws, [14, 42, 20, 45])
    ws.freeze_panes = "A2"


def write_programme_leaders_sheet(wb, pm_rows):
    ws = wb.create_sheet("Programme Leaders")
    headers = ["Staff Name", "Staff ID", "Role", "Hours", "Programme(s)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 22

    for i, (name, sid, role, hrs) in enumerate(pm_rows, 2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=sid)
        ws.cell(row=i, column=3, value=role)
        ws.cell(row=i, column=4, value=hrs)
        ws.cell(row=i, column=5, value="")
        style_row(ws, i, len(headers), alt=(i % 2 == 0))
        ws.row_dimensions[i].height = 18

    set_col_widths(ws, [28, 12, 32, 10, 45])
    ws.freeze_panes = "A2"


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Generate admin spreadsheet from Resources Planner")
    parser.add_argument("workbook", help="Resources Planner workbook (.xlsm)")
    parser.add_argument("-o", "--output", default="admin_sheet.xlsx",
                        help="Output .xlsx path (default: admin_sheet.xlsx)")
    args = parser.parse_args()

    print(f"Reading {args.workbook} …")
    wb_in = openpyxl.load_workbook(args.workbook, data_only=True)
    ws    = wb_in[SHEET]
    sec   = find_sections(ws)

    tutors  = build_module_tutors(ws, sec)
    periods = build_module_period_map(ws, sec)
    titles  = build_module_title_map(ws, sec)
    pm_rows = build_programme_leaders(ws, sec)

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    write_module_tutors_sheet(wb_out, tutors, periods, titles)
    write_programme_leaders_sheet(wb_out, pm_rows)

    out_path = Path(args.output)
    wb_out.save(out_path)
    print(f"Written: {out_path}")
    print(f"  Module Tutors:     {len(tutors)} modules")
    print(f"  Programme Leaders: {len(pm_rows)} rows")


if __name__ == "__main__":
    main()
